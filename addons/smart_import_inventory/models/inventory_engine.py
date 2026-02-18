import base64
import csv
import io
import json
import logging
from collections import defaultdict

from openpyxl import load_workbook

from odoo import _, fields, models

_logger = logging.getLogger(__name__)


class SmartImportInventoryEngine(models.AbstractModel):
    _name = 'smart.import.inventory.engine'
    _description = 'Smart Import Inventory Engine'

    PLAN_ORDER = [
        'uoms',
        'categories',
        'vendors',
        'locations_putaway',
        'products',
        'orderpoints',
        'boms',
    ]

    def _to_float(self, value):
        if value in (False, None, ''):
            return 0.0
        try:
            return float(value)
        except (TypeError, ValueError):
            return 0.0

    def _to_bool(self, value):
        if isinstance(value, bool):
            return value
        return str(value or '').strip().lower() in ('1', 'true', 'yes', 'y')

    def _normalized(self, value):
        return str(value or '').strip()

    def _uom_key(self, value):
        return self._normalized(value).lower()

    def _resolve_uom(self, uom_name, uom_cache, uom_alias_cache):
        if not uom_name:
            return False
        if uom_name in uom_cache:
            return uom_cache[uom_name]
        alias_uom = uom_alias_cache.get(self._uom_key(uom_name))
        if alias_uom:
            uom_cache[uom_name] = alias_uom
            return alias_uom
        return False

    def _build_uom_alias_cache(self, env):
        alias_cache = {}
        aliases_by_xmlid = {
            'uom.product_uom_unit': ['piece', 'pieces', 'unit', 'units', 'set', 'sets'],
            'uom.product_uom_kgm': ['kg', 'kilogram', 'kilograms'],
            'uom.product_uom_litre': ['l', 'liter', 'litre'],
            'uom.product_uom_millilitre': ['ml', 'milliliter', 'millilitre'],
            'uom.product_uom_inch': ['in', 'inch', 'inches'],
            'uom.product_uom_foot': ['ft', 'foot', 'feet'],
        }
        for xmlid, aliases in aliases_by_xmlid.items():
            uom = env.ref(xmlid, raise_if_not_found=False)
            if not uom:
                continue
            alias_cache[self._uom_key(uom.name)] = uom
            alias_cache[self._uom_key(uom.display_name)] = uom
            for alias in aliases:
                alias_cache[self._uom_key(alias)] = uom
        return alias_cache

    def _header_lookup_key(self, value):
        return self._normalized(value).lower()

    def _get_default_mapping_profile(self):
        company = self.env.company
        profile = self.env['smart.import.mapping.profile'].search([
            ('name', '=', 'Fulcrum Default'),
            ('active', '=', True),
            ('company_id', '=', company.id),
        ], limit=1)
        if not profile:
            profile = self.env['smart.import.mapping.profile'].search([
                ('name', '=', 'Fulcrum Default'),
                ('active', '=', True),
                ('company_id', '=', False),
            ], limit=1)
        if not profile:
            profile = self.env['smart.import.mapping.profile'].search([
                ('active', '=', True),
                "|", ('company_id', '=', company.id), ('company_id', '=', False),
            ], limit=1)
        if not profile:
            profile = self.env['smart.import.mapping.profile'].search([('active', '=', True)], limit=1)
        return profile

    def _build_header_map(self, headers, mapping_profile=False):
        profile = mapping_profile or self._get_default_mapping_profile()
        header_map = {}
        if not profile:
            return header_map
        keyword_map = {}
        for keyword in profile.keyword_ids.filtered('active'):
            key = self._header_lookup_key(keyword.source_key)
            if not key:
                continue
            current = keyword_map.get(key)
            if not current or keyword.priority < current.priority:
                keyword_map[key] = keyword
        for header in headers:
            if not header:
                continue
            match = keyword_map.get(self._header_lookup_key(header))
            if match:
                header_map[header] = match.canonical_key
        return header_map

    def apply_rules(self, profile, row, proposed_vals, explain_log_list, applies_to='product_template'):
        if not profile:
            return 0
        match_count = 0
        for rule in profile.rule_ids.filtered(lambda r: r.active and r.applies_to == applies_to):
            source_value = row.get(rule.when_key)
            if not rule.predicate_matches(source_value):
                continue
            if rule.set_purchase_ok:
                proposed_vals['purchase_ok'] = True
            if rule.add_route_buy:
                proposed_vals['add_route_buy'] = True
            if rule.add_route_manufacture:
                proposed_vals['add_route_manufacture'] = True
            explanation = _(
                "Applied rule '%(rule)s' because %(key)s='%(value)s' %(operator)s '%(expected)s'"
            ) % {
                'rule': rule.name,
                'key': rule.when_key,
                'value': self._normalized(source_value),
                'operator': rule.operator,
                'expected': rule.when_value,
            }
            explain_log_list.append(explanation)
            match_count += 1
        return match_count

    def _apply_session_overrides(self, proposed_vals, options):
        if options.get('force_purchase_ok'):
            proposed_vals['purchase_ok'] = True
        if options.get('force_buy_route'):
            proposed_vals['add_route_buy'] = True
        if options.get('force_manufacture_route'):
            proposed_vals['add_route_manufacture'] = True

    def _build_product_proposal(self, row, profile=False, options=False, explain_log_list=False):
        explain_log_list = explain_log_list if explain_log_list is not False else []
        proposed_vals = {
            'purchase_ok': row['buy_or_make'] == 'Buy' or bool(row['vendor_name']),
            'add_route_buy': row['buy_or_make'] == 'Buy' or bool(row['vendor_name']),
            'add_route_manufacture': row['buy_or_make'] == 'Make',
        }
        rule_matches = self.apply_rules(profile, row, proposed_vals, explain_log_list)
        self._apply_session_overrides(proposed_vals, options or {})
        return proposed_vals, rule_matches

    def parse_items_xlsx(self, attachment, mapping_profile=False):
        data = base64.b64decode(attachment.datas)
        workbook = load_workbook(io.BytesIO(data), data_only=True)
        sheet = workbook['Filtered Items'] if 'Filtered Items' in workbook.sheetnames else workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [self._normalized(head) for head in rows[0]]
        header_map = self._build_header_map(headers, mapping_profile=mapping_profile)
        parsed_rows = []
        for row in rows[1:]:
            if not any(row):
                continue
            vals = {headers[idx]: row[idx] if idx < len(row) else False for idx in range(len(headers))}
            canonical_vals = {}
            extra_columns = {}
            for header, value in vals.items():
                canonical_key = header_map.get(header)
                if canonical_key:
                    canonical_vals[canonical_key] = value
                else:
                    extra_columns[header] = value

            parsed_rows.append({
                'default_code': self._normalized(canonical_vals.get('default_code')),
                'name': self._normalized(canonical_vals.get('name')),
                'tags': self._normalized(canonical_vals.get('tags')),
                'buy_or_make': self._normalized(canonical_vals.get('buy_or_make')),
                'min_stock': self._to_float(canonical_vals.get('min_stock')),
                'min_production_qty': self._to_float(canonical_vals.get('min_production_qty')),
                'uom_name': self._normalized(canonical_vals.get('uom_name')),
                'category': self._normalized(canonical_vals.get('category')),
                'sell_ok': self._to_bool(canonical_vals.get('sell_ok')),
                'default_location': self._normalized(canonical_vals.get('default_location')),
                'vendor_name': self._normalized(canonical_vals.get('vendor_name')),
                'vendor_price': self._to_float(canonical_vals.get('vendor_price')),
                'vendor_moq': self._to_float(canonical_vals.get('vendor_moq')),
                'vendor_uom': self._normalized(canonical_vals.get('vendor_uom')),
                'raw': vals,
                '_extra_columns': extra_columns,
                '_explain': [],
            })
        return parsed_rows

    def parse_bom_xlsx(self, attachment):
        data = base64.b64decode(attachment.datas)
        workbook = load_workbook(io.BytesIO(data), data_only=True)
        sheet = workbook['Sheet1'] if 'Sheet1' in workbook.sheetnames else workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [self._normalized(head) for head in rows[0]]
        parsed_rows = []
        for row in rows[1:]:
            if not any(row):
                continue
            vals = {headers[idx]: row[idx] if idx < len(row) else False for idx in range(len(headers))}
            parsed_rows.append({
                'parent_number': self._normalized(vals.get('Parent Number')),
                'parent_description': self._normalized(vals.get('Parent Description')),
                'child_number': self._normalized(vals.get('Child Number')),
                'child_description': self._normalized(vals.get('Child Description')),
                'units_required': self._to_float(vals.get('Units Required')) or 1.0,
            })
        return parsed_rows

    def build_plan(self, items_rows, bom_rows, mapping_profile=False, options=False):
        options = options or {}
        unique_uoms = sorted({r['uom_name'] for r in items_rows if r['uom_name']})
        unique_categories = sorted({r['category'] for r in items_rows if r['category']})
        unique_vendors = sorted({r['vendor_name'] for r in items_rows if r['vendor_name']})
        unique_locations = sorted({r['default_location'] for r in items_rows if r['default_location']})
        product_codes = {r['default_code'] for r in items_rows if r['default_code']}

        bom_parent_codes = {r['parent_number'] for r in bom_rows if r['parent_number']}
        bom_child_codes = {r['child_number'] for r in bom_rows if r['child_number']}
        missing_bom_products = sorted((bom_parent_codes | bom_child_codes) - product_codes)

        uom_cache = {u.name: u for u in self.env['uom.uom'].search([])}
        uom_alias_cache = self._build_uom_alias_cache(self.env)
        unknown_uoms = []
        for row in items_rows:
            for uom_name in (row['uom_name'], row['vendor_uom']):
                if uom_name and not self._resolve_uom(uom_name, uom_cache, uom_alias_cache):
                    unknown_uoms.append(uom_name)
        unknown_uoms = sorted(set(unknown_uoms))

        grouped_boms = defaultdict(list)
        for row in bom_rows:
            if row['parent_number']:
                grouped_boms[row['parent_number']].append(row)

        plan = {
            'steps': self.PLAN_ORDER,
            'counts': {
                'unique_uoms': len(unique_uoms),
                'unique_categories': len(unique_categories),
                'unique_vendors': len(unique_vendors),
                'unique_locations': len(unique_locations),
                'product_rows_count': len(items_rows),
                'orderpoints_count': len([r for r in items_rows if r['min_stock'] > 0]),
                'boms_count': len(grouped_boms),
                'bom_lines_count': len(bom_rows),
            },
            'mapped_fields': {
                'product.template': ['default_code', 'name', 'categ_id', 'uom_id', 'uom_po_id', 'sale_ok', 'purchase_ok', 'type', 'route_ids'],
                'product.supplierinfo': ['partner_id', 'price', 'min_qty'],
                'stock.warehouse.orderpoint': ['product_id', 'location_id', 'product_min_qty', 'product_max_qty', 'qty_multiple'],
                'mrp.bom': ['product_tmpl_id'],
                'mrp.bom.line': ['product_id', 'product_qty', 'product_uom_id'],
            },
            'issues': {
                'unknown_uoms': unknown_uoms,
                'missing_bom_products': missing_bom_products,
                'rows_without_product_number': [idx + 2 for idx, row in enumerate(items_rows) if not row['default_code']],
            },
            'route_preview_counts': {
                'count_need_buy': 0,
                'count_need_manufacture': 0,
                'count_need_both': 0,
            },
            'sample_previews': [],
            'proposed_defaults': {
                'product_type_policy': _('Default product type is Consumable for imported products.'),
                'route_policy_summary': _('Routes are proposed from Buy/Make values, vendor availability, mapping rules, and session override toggles.'),
                'purchase_ok_policy_summary': _('Purchase is enabled from Buy/Make values, vendor availability, mapping rules, and can be forced globally from the wizard.'),
                'orderpoint_policy_summary': _('Orderpoints are prepared when minimum stock is set. Maximum quantity follows the selected max policy.'),
            },
        }

        for row in items_rows:
            proposed_vals, _rule_matches = self._build_product_proposal(
                row,
                profile=mapping_profile,
                options=options,
                explain_log_list=[],
            )
            if proposed_vals['add_route_buy']:
                plan['route_preview_counts']['count_need_buy'] += 1
            if proposed_vals['add_route_manufacture']:
                plan['route_preview_counts']['count_need_manufacture'] += 1
            if proposed_vals['add_route_buy'] and proposed_vals['add_route_manufacture']:
                plan['route_preview_counts']['count_need_both'] += 1

            if len(plan['sample_previews']) < 10:
                proposed_routes = []
                if proposed_vals['add_route_buy']:
                    proposed_routes.append('Buy')
                if proposed_vals['add_route_manufacture']:
                    proposed_routes.append('Manufacture')
                plan['sample_previews'].append({
                    'default_code': row['default_code'],
                    'name': row['name'],
                    'buy_or_make': row['buy_or_make'],
                    'purchase_ok': proposed_vals['purchase_ok'],
                    'routes': proposed_routes,
                })

        return plan

    def execute(self, session, items_rows, bom_rows, options, dry_run=False):
        errors = []
        stats = defaultdict(int)
        issue_bucket = defaultdict(list)
        env = self.env

        company = env.company
        warehouse = env['stock.warehouse'].search([('company_id', '=', company.id)], limit=1)
        stock_location = warehouse.lot_stock_id if warehouse else env.ref('stock.stock_location_stock')

        def add_log(message, level='INFO'):
            _logger.info('Fulcrum import %s: %s', session.name, message)
            session.append_log(level, message)

        add_log(_('Starting %s execution') % ('dry-run' if dry_run else 'import'))
        mapping_profile = self.env['smart.import.mapping.profile'].browse(options.get('mapping_profile_id')) if options.get('mapping_profile_id') else self._get_default_mapping_profile()

        # Caches
        uom_cache = {u.name: u for u in env['uom.uom'].search([])}
        uom_alias_cache = self._build_uom_alias_cache(env)
        categ_cache = {c.name: c for c in env['product.category'].search([])}
        partner_cache = {p.name: p for p in env['res.partner'].search([('supplier_rank', '>', 0)])}
        product_tmpl_cache = {p.default_code: p for p in env['product.template'].search([('default_code', '!=', False)])}
        product_variant_cache = {p.default_code: p.product_variant_id for p in product_tmpl_cache.values() if p.product_variant_id}
        location_cache = {l.complete_name: l for l in env['stock.location'].search([('company_id', '=', company.id), ('usage', '=', 'internal')])}

        buy_route = env['stock.route'].search([('name', '=', 'Buy'), ('company_id', 'in', [False, company.id])], limit=1)
        manufacture_route = env['stock.route'].search([('name', '=', 'Manufacture'), ('company_id', 'in', [False, company.id])], limit=1)

        # UoMs
        ref_uom = env.ref('uom.product_uom_unit')
        for uom_name in sorted({r['uom_name'] for r in items_rows if r['uom_name']} | {r['vendor_uom'] for r in items_rows if r['vendor_uom']}):
            if self._resolve_uom(uom_name, uom_cache, uom_alias_cache):
                continue
            if not options.get('auto_create_unknown_uom'):
                issue_bucket['unknown_uoms'].append(uom_name)
                continue
            if dry_run:
                stats['uom_would_create'] += 1
                # Simulate availability for subsequent validation paths (products/vendors/BOM).
                uom_cache[uom_name] = ref_uom
                continue
            created = env['uom.uom'].create({
                'name': uom_name,
                'uom_type': 'smaller',
                'factor_inv': 1.0,
                'category_id': ref_uom.category_id.id,
                'rounding': ref_uom.rounding,
            })
            uom_cache[uom_name] = created
            stats['uom_created'] += 1
            add_log(_('Created UoM %s') % uom_name)

        # Categories
        for category_name in sorted({r['category'] for r in items_rows if r['category']}):
            if category_name in categ_cache:
                continue
            if dry_run:
                stats['category_would_create'] += 1
                categ_cache[category_name] = env.ref('product.product_category_all')
                continue
            created = env['product.category'].create({'name': category_name})
            categ_cache[category_name] = created
            stats['category_created'] += 1
            add_log(_('Created category %s') % category_name)

        # Vendors
        for vendor_name in sorted({r['vendor_name'] for r in items_rows if r['vendor_name']}):
            if vendor_name in partner_cache:
                continue
            if dry_run:
                stats['vendor_would_create'] += 1
                partner_cache[vendor_name] = env.user.partner_id
                continue
            created = env['res.partner'].create({'name': vendor_name, 'supplier_rank': 1, 'company_type': 'company'})
            partner_cache[vendor_name] = created
            stats['vendor_created'] += 1
            add_log(_('Created vendor %s') % vendor_name)

        # Locations + putaway
        if options.get('create_locations_putaway'):
            for loc_code in sorted({r['default_location'] for r in items_rows if r['default_location']}):
                full_name = f'{stock_location.complete_name}/Fulcrum/{loc_code}'
                existing = location_cache.get(full_name)
                if existing:
                    continue
                if dry_run:
                    stats['location_would_create'] += 1
                    location_cache[full_name] = stock_location
                    continue
                parent = env['stock.location'].search([
                    ('name', '=', 'Fulcrum'),
                    ('location_id', '=', stock_location.id),
                    ('usage', '=', 'internal'),
                    ('company_id', '=', company.id),
                ], limit=1)
                if not parent:
                    parent = env['stock.location'].create({
                        'name': 'Fulcrum',
                        'location_id': stock_location.id,
                        'usage': 'internal',
                        'company_id': company.id,
                    })
                created = env['stock.location'].create({
                    'name': loc_code,
                    'location_id': parent.id,
                    'usage': 'internal',
                    'company_id': company.id,
                })
                location_cache[created.complete_name] = created
                stats['location_created'] += 1
                add_log(_('Created location %s') % created.complete_name)

        # Products + supplierinfo + orderpoints
        product_rows_by_code = {}
        for row in items_rows:
            if not row['default_code']:
                issue_bucket['rows_without_product_number'].append(row)
                continue
            product_rows_by_code[row['default_code']] = row

        for code, row in product_rows_by_code.items():
            uom_name = row['uom_name']
            uom = self._resolve_uom(uom_name, uom_cache, uom_alias_cache) if uom_name else False
            if uom_name and not uom:
                issue_bucket['unknown_uoms'].append(uom_name)
                errors.append({'type': 'product', 'key': code, 'error': f'Unknown UoM {uom_name}'})
                continue

            category = categ_cache.get(row['category'])
            row_explain = row.setdefault('_explain', [])
            proposed_vals, rule_matches = self._build_product_proposal(
                row,
                profile=mapping_profile,
                options=options,
                explain_log_list=row_explain,
            )
            if rule_matches:
                stats['rules_applied'] += rule_matches
                for message in row_explain[-rule_matches:]:
                    add_log(_('%(code)s: %(message)s') % {'code': code, 'message': message})

            vals = {
                'default_code': code,
                'name': row['name'] or code,
                'sale_ok': row['sell_ok'],
                'purchase_ok': proposed_vals['purchase_ok'],
                'type': 'consu',
            }
            if category:
                vals['categ_id'] = category.id
            if uom:
                vals['uom_id'] = uom.id
                vals['uom_po_id'] = uom.id

            route_ids = []
            if proposed_vals['add_route_manufacture'] and manufacture_route:
                route_ids.append(manufacture_route.id)
            if proposed_vals['add_route_buy'] and buy_route:
                route_ids.append(buy_route.id)
            if route_ids:
                vals['route_ids'] = [(6, 0, route_ids)]

            tmpl = product_tmpl_cache.get(code)
            if tmpl:
                if dry_run:
                    stats['product_would_update'] += 1
                else:
                    tmpl.write(vals)
                    stats['product_updated'] += 1
                    add_log(_('Updated product %s') % code)
            else:
                if dry_run:
                    stats['product_would_create'] += 1
                    tmpl = env['product.template']
                    product_tmpl_cache[code] = tmpl
                    product_variant_cache[code] = env['product.product']
                else:
                    tmpl = env['product.template'].create(vals)
                    stats['product_created'] += 1
                    add_log(_('Created product %s') % code)
                    product_tmpl_cache[code] = tmpl
                    product_variant_cache[code] = tmpl.product_variant_id

            partner = partner_cache.get(row['vendor_name']) if row['vendor_name'] else False
            if partner and not dry_run and tmpl:
                supp_vals = {
                    'partner_id': partner.id,
                    'price': row['vendor_price'],
                    'min_qty': row['vendor_moq'] or 0.0,
                    'product_tmpl_id': tmpl.id,
                }
                vendor_uom_name = row['vendor_uom']
                vendor_uom = self._resolve_uom(vendor_uom_name, uom_cache, uom_alias_cache) if vendor_uom_name else False
                if vendor_uom:
                    supp_vals['product_uom'] = vendor_uom.id
                elif vendor_uom_name:
                    issue_bucket['unknown_vendor_uom'].append(vendor_uom_name)
                supplierinfo = env['product.supplierinfo'].search([
                    ('partner_id', '=', partner.id),
                    ('product_tmpl_id', '=', tmpl.id),
                ], limit=1)
                if supplierinfo:
                    supplierinfo.write(supp_vals)
                    stats['supplierinfo_updated'] += 1
                else:
                    env['product.supplierinfo'].create(supp_vals)
                    stats['supplierinfo_created'] += 1

            if row['min_stock'] > 0:
                product_variant = product_variant_cache.get(code)
                if not product_variant and code not in product_variant_cache:
                    continue

                max_qty = row['min_stock']
                if options.get('orderpoint_max_policy') == 'double_min':
                    max_qty = row['min_stock'] * 2

                if dry_run and not product_variant:
                    stats['orderpoint_would_create'] += 1
                    continue

                op_vals = {
                    'product_id': product_variant.id,
                    'location_id': stock_location.id,
                    'product_min_qty': row['min_stock'],
                    'product_max_qty': max_qty,
                    'qty_multiple': row['min_production_qty'] if row['min_production_qty'] > 0 else 1.0,
                    'company_id': company.id,
                }
                op = env['stock.warehouse.orderpoint'].search([
                    ('product_id', '=', product_variant.id),
                    ('location_id', '=', stock_location.id),
                    ('company_id', '=', company.id),
                ], limit=1)
                if op:
                    if dry_run:
                        stats['orderpoint_would_update'] += 1
                    else:
                        op.write(op_vals)
                        stats['orderpoint_updated'] += 1
                else:
                    if dry_run:
                        stats['orderpoint_would_create'] += 1
                    else:
                        env['stock.warehouse.orderpoint'].create(op_vals)
                        stats['orderpoint_created'] += 1

            if options.get('create_locations_putaway') and row['default_location'] and product_variant_cache.get(code):
                full_name = f'{stock_location.complete_name}/Fulcrum/{row["default_location"]}'
                dst_location = location_cache.get(full_name)
                if dst_location and not dry_run:
                    putaway = env['stock.putaway.rule'].search([
                        ('product_id', '=', product_variant_cache[code].id),
                        ('location_in_id', '=', stock_location.id),
                        ('location_out_id', '=', dst_location.id),
                        ('company_id', '=', company.id),
                    ], limit=1)
                    if not putaway:
                        env['stock.putaway.rule'].create({
                            'product_id': product_variant_cache[code].id,
                            'location_in_id': stock_location.id,
                            'location_out_id': dst_location.id,
                            'company_id': company.id,
                        })
                        stats['putaway_created'] += 1

        # BOMs
        boms_by_parent = defaultdict(list)
        for row in bom_rows:
            if row['parent_number']:
                boms_by_parent[row['parent_number']].append(row)

        for parent_code, lines in boms_by_parent.items():
            if parent_code not in product_tmpl_cache:
                issue_bucket['missing_bom_parents'].append(parent_code)
                errors.append({'type': 'bom', 'key': parent_code, 'error': 'Missing parent product'})
                continue
            parent_tmpl = product_tmpl_cache[parent_code]
            if dry_run:
                stats['bom_would_process'] += 1
                continue

            bom = env['mrp.bom'].search([
                ('product_tmpl_id', '=', parent_tmpl.id),
                ('company_id', 'in', [False, company.id]),
                ('type', '=', 'normal'),
            ], limit=1)
            if not bom:
                bom = env['mrp.bom'].create({
                    'product_tmpl_id': parent_tmpl.id,
                    'product_qty': 1.0,
                    'type': 'normal',
                    'company_id': company.id,
                })
                stats['bom_created'] += 1
            else:
                stats['bom_updated'] += 1
                bom.bom_line_ids.unlink()

            for line in lines:
                child_code = line['child_number']
                child_product = product_variant_cache.get(child_code)
                if not child_product:
                    if options.get('create_placeholder_missing_bom_children') and child_code:
                        child_tmpl = env['product.template'].create({
                            'default_code': child_code,
                            'name': line['child_description'] or child_code,
                            'type': 'consu',
                            'sale_ok': False,
                            'purchase_ok': False,
                        })
                        child_product = child_tmpl.product_variant_id
                        product_tmpl_cache[child_code] = child_tmpl
                        product_variant_cache[child_code] = child_product
                        stats['placeholder_products_created'] += 1
                        add_log(_('Created placeholder BOM child %s') % child_code, level='WARNING')
                    else:
                        issue_bucket['missing_bom_children'].append(child_code)
                        errors.append({'type': 'bom_line', 'key': f'{parent_code}->{child_code}', 'error': 'Missing child product'})
                        continue
                env['mrp.bom.line'].create({
                    'bom_id': bom.id,
                    'product_id': child_product.id,
                    'product_qty': line['units_required'] or 1.0,
                    'product_uom_id': child_product.uom_id.id,
                })
                stats['bom_line_created'] += 1

        issue_payload = {k: sorted(set(filter(None, v))) for k, v in issue_bucket.items()}
        issue_payload['rule_applied_count'] = stats.get('rules_applied', 0)
        session_vals = {
            'stats_json': json.dumps(dict(stats), indent=2, sort_keys=True),
            'issues_json': json.dumps(issue_payload, indent=2, sort_keys=True),
            'state': 'planned' if dry_run else ('done' if not errors else 'error'),
        }
        session.write(session_vals)

        if errors:
            attachment = self._build_error_csv_attachment(session, errors)
            session.error_csv_attachment_id = attachment.id
            add_log(_('Generated error report with %s rows') % len(errors), level='WARNING')

        add_log(_('Execution finished with %s error rows') % len(errors), level='INFO')
        return {
            'stats': dict(stats),
            'issues': issue_payload,
            'errors': errors,
        }

    def _build_error_csv_attachment(self, session, errors):
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=['type', 'key', 'error'])
        writer.writeheader()
        for error in errors:
            writer.writerow(error)
        data = base64.b64encode(output.getvalue().encode('utf-8'))
        return self.env['ir.attachment'].create({
            'name': f'{session.name}_errors.csv',
            'type': 'binary',
            'datas': data,
            'mimetype': 'text/csv',
            'res_model': session._name,
            'res_id': session.id,
        })
