import base64
import json
import logging
from collections import defaultdict

from openpyxl import load_workbook

from odoo import _, models

_logger = logging.getLogger(__name__)


class SmartImportPackInventoryFulcrum(models.AbstractModel):
    _name = 'smart.import.pack.inventory_fulcrum'
    _inherit = 'smart.import.pack'
    _description = 'Smart Import Pack Inventory Fulcrum'

    _pack_code = 'inventory_fulcrum'
    _pack_name = 'Inventory (Fulcrum)'

    PLAN_ORDER = ['uoms', 'categories', 'vendors', 'locations_putaway', 'products', 'orderpoints', 'boms']

    def _normalized(self, value):
        return str(value or '').strip()

    def _to_float(self, value):
        if value in (False, None, ''):
            return 0.0
        try:
            return float(value)
        except (TypeError, ValueError):
            return 0.0

    def _to_bool(self, value):
        if isinstance(value, bool):
            return value
        return str(value or '').strip().lower() in ('1', 'true', 'yes', 'y')

    def _header_lookup_key(self, value):
        return self._normalized(value).lower()

    def _build_header_map(self, headers, mapping_profile=False):
        profile = mapping_profile
        header_map = {}
        if not profile:
            return header_map
        keyword_map = {}
        for keyword in profile.keyword_ids.filtered('active'):
            key = self._header_lookup_key(keyword.source_key)
            if key and (key not in keyword_map or keyword.priority < keyword_map[key].priority):
                keyword_map[key] = keyword
        for header in headers:
            match = keyword_map.get(self._header_lookup_key(header))
            if match:
                header_map[header] = match.canonical_key
        return header_map

    def _detect_file_kind(self, headers):
        header_keys = {self._header_lookup_key(h) for h in headers if h}
        if {'parent number', 'child number', 'units required'}.issubset(header_keys):
            return 'bom', [{'model': 'mrp.bom', 'confidence': 0.99}, {'model': 'mrp.bom.line', 'confidence': 0.99}]
        if {'number', 'description'}.issubset(header_keys):
            return 'items', [
                {'model': 'product.template', 'confidence': 0.95},
                {'model': 'product.supplierinfo', 'confidence': 0.8},
                {'model': 'stock.warehouse.orderpoint', 'confidence': 0.7},
            ]
        return 'unknown', []

    def detect(self, session):
        summary = {'items': 0, 'bom': 0, 'unknown': 0}
        for line in session.file_line_ids:
            headers = json.loads(line.header_row_json or '[]')
            kind, models_payload = self._detect_file_kind(headers)
            line.write({'kind': kind, 'detected_models_json': json.dumps(models_payload)})
            summary[kind] += 1
        session.append_log('INFO', _('Inventory pack detect complete.'))
        return summary

    def parse_items_xlsx(self, attachment, mapping_profile=False):
        import io
        data = base64.b64decode(attachment.datas)
        workbook = load_workbook(io.BytesIO(data), data_only=True)
        sheet = workbook['Sheet1'] if 'Sheet1' in workbook.sheetnames else workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [self._normalized(head) for head in rows[0]]
        header_map = self._build_header_map(headers, mapping_profile=mapping_profile)
        parsed_rows = []
        for row in rows[1:]:
            if not any(row):
                continue
            vals = {headers[idx]: row[idx] if idx < len(row) else False for idx in range(len(headers))}
            canonical_vals = {}
            for header, canonical_key in header_map.items():
                canonical_vals[canonical_key] = vals.get(header)
            parsed_rows.append({
                'default_code': self._normalized(canonical_vals.get('default_code')),
                'name': self._normalized(canonical_vals.get('name')),
                'category': self._normalized(canonical_vals.get('category')),
                'uom_name': self._normalized(canonical_vals.get('uom_name')),
                'vendor_name': self._normalized(canonical_vals.get('vendor_name')),
                'vendor_price': self._to_float(canonical_vals.get('vendor_price')),
                'vendor_moq': self._to_float(canonical_vals.get('vendor_moq')),
                'vendor_uom': self._normalized(canonical_vals.get('vendor_uom')),
                'default_location': self._normalized(canonical_vals.get('default_location')),
                'min_stock': self._to_float(canonical_vals.get('min_stock')),
                'min_production_qty': self._to_float(canonical_vals.get('min_production_qty')),
                'sell_ok': self._to_bool(canonical_vals.get('sell_ok')),
                'buy_or_make': self._normalized(canonical_vals.get('buy_or_make')),
            })
        return parsed_rows

    def parse_bom_xlsx(self, attachment):
        import io
        data = base64.b64decode(attachment.datas)
        workbook = load_workbook(io.BytesIO(data), data_only=True)
        sheet = workbook['Sheet1'] if 'Sheet1' in workbook.sheetnames else workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [self._normalized(head) for head in rows[0]]
        parsed_rows = []
        for row in rows[1:]:
            if not any(row):
                continue
            vals = {headers[idx]: row[idx] if idx < len(row) else False for idx in range(len(headers))}
            parsed_rows.append({
                'parent_number': self._normalized(vals.get('Parent Number')),
                'parent_description': self._normalized(vals.get('Parent Description')),
                'child_number': self._normalized(vals.get('Child Number')),
                'child_description': self._normalized(vals.get('Child Description')),
                'units_required': self._to_float(vals.get('Units Required')) or 1.0,
            })
        return parsed_rows

    def _get_rows_from_session(self, session):
        profile = session.mapping_profile_id
        items_rows, bom_rows = [], []
        for line in session.file_line_ids:
            if line.kind == 'items':
                items_rows.extend(self.parse_items_xlsx(line.attachment_id, mapping_profile=profile))
            elif line.kind == 'bom':
                bom_rows.extend(self.parse_bom_xlsx(line.attachment_id))
        return items_rows, bom_rows

    def build_plan(self, session, dry_run=True):
        items_rows, bom_rows = self._get_rows_from_session(session)
        product_codes = {r['default_code'] for r in items_rows if r['default_code']}
        grouped_boms = defaultdict(list)
        for row in bom_rows:
            if row['parent_number']:
                grouped_boms[row['parent_number']].append(row)
        missing_bom_products = sorted(({r['parent_number'] for r in bom_rows if r['parent_number']} | {r['child_number'] for r in bom_rows if r['child_number']}) - product_codes)
        stats = {
            'steps': self.PLAN_ORDER,
            'counts': {
                'unique_uoms': len({r['uom_name'] for r in items_rows if r['uom_name']}),
                'unique_categories': len({r['category'] for r in items_rows if r['category']}),
                'unique_vendors': len({r['vendor_name'] for r in items_rows if r['vendor_name']}),
                'unique_locations': len({r['default_location'] for r in items_rows if r['default_location']}),
                'products': len(items_rows),
                'orderpoints': len([r for r in items_rows if r['min_stock'] > 0]),
                'boms': len(grouped_boms),
                'bom_lines': len(bom_rows),
            },
            'bom_parent_codes': sorted(grouped_boms.keys()),
            'route_preview_counts': {'count_need_buy': 0, 'count_need_manufacture': 0, 'count_need_both': 0},
        }
        issues = {'missing_bom_products': missing_bom_products}
        session.append_log('INFO', _('Step 1 UoM: %s unique') % stats['counts']['unique_uoms'])
        session.append_log('INFO', _('Step 2 Categories: %s') % stats['counts']['unique_categories'])
        session.append_log('INFO', _('Step 3 Products: %s') % stats['counts']['products'])
        session.append_log('INFO', _('Step 4 BOMs: %s parents / %s lines') % (stats['counts']['boms'], stats['counts']['bom_lines']))
        return {'stats': stats, 'issues': issues, 'dry_run': dry_run}

    def execute(self, session, dry_run=False):
        items_rows, bom_rows = self._get_rows_from_session(session)
        legacy = self.env['smart.import.inventory.engine']
        options = json.loads(session.options_json or '{}')
        options.update({'mapping_profile_id': session.mapping_profile_id.id})
        session.append_log('INFO', _('Delegating execution to legacy engine compatibility path.'))
        result = legacy.execute(session, items_rows, bom_rows, options, dry_run=dry_run)
        if result.get('errors'):
            session.make_error_csv(result['errors'])
        return result
