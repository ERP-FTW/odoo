import base64
import io

from openpyxl import load_workbook

from odoo import models


class SmartImportXlsxHelper(models.AbstractModel):
    _name = 'smart.import.xlsx.helper'
    _description = 'Smart Import XLSX Helper'

    def _workbook_from_attachment(self, attachment):
        data = base64.b64decode(attachment.datas or b'')
        return load_workbook(io.BytesIO(data), data_only=True)

    def get_sheet_names(self, attachment):
        workbook = self._workbook_from_attachment(attachment)
        return workbook.sheetnames

    def read_headers(self, attachment, sheet_name=False, header_row_guess=1):
        workbook = self._workbook_from_attachment(attachment)
        sheet = workbook[sheet_name] if sheet_name and sheet_name in workbook.sheetnames else workbook.active
        row = list(sheet.iter_rows(min_row=header_row_guess, max_row=header_row_guess, values_only=True))
        if not row:
            return []
        return [str(h).strip() if h not in (None, False) else '' for h in row[0]]

    def read_preview_rows(self, attachment, sheet_name=False, n=10, header_row=1):
        workbook = self._workbook_from_attachment(attachment)
        sheet = workbook[sheet_name] if sheet_name and sheet_name in workbook.sheetnames else workbook.active
        rows = []
        for values in sheet.iter_rows(min_row=header_row + 1, max_row=header_row + n, values_only=True):
            vals = [str(v).strip() if v not in (None, False) else '' for v in values]
            if any(vals):
                rows.append(vals)
        return rows
